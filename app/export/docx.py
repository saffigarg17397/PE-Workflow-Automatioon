"""DOCX export.

Deal teams circulate Word documents, so the memo has to leave the tool in a form
that survives being emailed and marked up. Citations render inline as [p.N] the
same as the Markdown export — a memo that loses its provenance on export would
defeat the point.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.models.flags import DetectionSource
from app.models.memo import MemoRun

NAVY = RGBColor(0x1A, 0x2F, 0x4B)
GREY = RGBColor(0x6B, 0x72, 0x80)

_REC_LABEL = {
    "advance": "ADVANCE",
    "more_info": "MORE INFORMATION REQUIRED",
    "pass": "PASS",
}


def render(run: MemoRun) -> bytes:
    m = run.memo
    p = run.profile
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_heading("Investment Committee Screening Memo", level=0)
    for r in title.runs:
        r.font.color.rgb = NAVY
    doc.add_heading(m.company_name, level=1)

    meta = doc.add_paragraph()
    mr = meta.add_run(
        f"Screened against {m.thesis_name} · {run.source_file} ({run.page_count}pp) · "
        f"{m.generated_at:%Y-%m-%d %H:%M} UTC · run {run.run_id}"
    )
    mr.font.size = Pt(8.5)
    mr.font.color.rgb = GREY

    disc = doc.add_paragraph()
    dr = disc.add_run(
        "Machine-generated first draft. Figures are cited to source pages; uncited "
        "prose is model-generated and unverified. Not an investment recommendation."
    )
    dr.italic = True
    dr.font.size = Pt(8.5)
    dr.font.color.rgb = GREY

    # 1
    doc.add_heading(f"1. Recommendation: {_REC_LABEL[m.recommendation.value]}", level=2)
    doc.add_paragraph(m.recommendation_rationale)

    # 2
    doc.add_heading("2. Business Overview", level=2)
    doc.add_paragraph(m.business_overview or "Not generated.")

    # 3
    doc.add_heading("3. Financial Summary", level=2)
    years = sorted(p.financials.value or [], key=lambda y: y.year)
    if years:
        t = doc.add_table(rows=1, cols=len(years) + 1)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].text = "($M)"
        for i, y in enumerate(years, 1):
            hdr[i].text = str(y.year)
        for label, attr in (
            ("Revenue", "revenue"),
            ("Gross profit", "gross_profit"),
            ("Reported EBITDA", "reported_ebitda"),
            ("Adjusted EBITDA", "adjusted_ebitda"),
        ):
            vals = [getattr(y, attr) for y in years]
            if all(v is None for v in vals):
                continue
            row = t.add_row().cells
            row[0].text = label
            for i, v in enumerate(vals, 1):
                row[i].text = f"{v:,.2f}" if v is not None else "n/a"
        margins = [y.ebitda_margin for y in years]
        if any(x is not None for x in margins):
            row = t.add_row().cells
            row[0].text = "Adj. EBITDA margin"
            for i, v in enumerate(margins, 1):
                row[i].text = f"{v:.1f}%" if v is not None else "n/a"
        doc.add_paragraph()

    derived = []
    if p.revenue_cagr is not None:
        derived.append(f"Revenue CAGR {p.revenue_cagr:.1f}%")
    if p.addback_pct_of_ebitda is not None:
        derived.append(f"add-backs {p.addback_pct_of_ebitda:.1f}% of adjusted EBITDA")
    if p.asking_multiple.value is not None:
        cites = "".join(c.render() for c in p.asking_multiple.unique_citations)
        derived.append(f"asking multiple {p.asking_multiple.value:.1f}x{cites}")
    if derived:
        doc.add_paragraph("Derived: " + " · ".join(derived) + ".")
    doc.add_paragraph(m.financial_summary or "")

    # 4
    fit = m.thesis_fit
    doc.add_heading(f"4. Thesis Fit — {fit.score:.0f}% ({fit.thesis_name})", level=2)
    cov = doc.add_paragraph()
    cr = cov.add_run(f"Coverage: {fit.coverage:.0f}% of criteria evaluable from this document.")
    cr.italic = True
    cr.font.size = Pt(9)

    ft = doc.add_table(rows=1, cols=4)
    ft.style = "Light Grid Accent 1"
    for i, header in enumerate(("Criterion", "Result", "Actual", "Target")):
        ft.rows[0].cells[i].text = header
    for c in fit.criteria:
        mark = {True: "pass", False: "FAIL", None: "n/a"}[c.passed]
        if c.is_dealbreaker and c.passed is False:
            mark = "FAIL — DEALBREAKER"
        row = ft.add_row().cells
        row[0].text = c.name
        row[1].text = mark
        row[2].text = c.actual
        row[3].text = c.target
    doc.add_paragraph()
    doc.add_paragraph(m.thesis_commentary or "")

    # 5
    doc.add_heading("5. Key Risks", level=2)
    if not m.red_flags:
        doc.add_paragraph("No red flags identified.")
    else:
        for source, label in (
            (DetectionSource.RULE, "Detected by rules (deterministic, reproducible)"),
            (
                DetectionSource.MODEL,
                "Identified in review (qualitative — verify before relying on)",
            ),
        ):
            group = [f for f in m.flags_by_severity if f.source == source]
            if not group:
                continue
            group_heading = doc.add_paragraph()
            group_heading.add_run(label).bold = True
            for flag in group:
                para = doc.add_paragraph(style="List Bullet")
                para.add_run(f"[{flag.severity.value.upper()}] ").bold = True
                cites = "".join(c.render() for c in flag.unique_citations)
                para.add_run(f"{flag.title}{cites} — {flag.detail}")

    # 6
    doc.add_heading("6. Diligence Priorities", level=2)
    for d in m.diligence_priorities or ["None generated."]:
        doc.add_paragraph(d, style="List Number")

    # 7
    doc.add_heading("7. Sources and Confidence", level=2)
    doc.add_paragraph(m.confidence_note)
    if m.unreviewed_fields:
        pending_heading = doc.add_paragraph()
        pending_heading.add_run("Pending human review:").bold = True
        for field_name in m.unreviewed_fields:
            cited = getattr(p, field_name)
            doc.add_paragraph(
                f"{field_name} — {cited.note or 'confidence: ' + cited.confidence.value}",
                style="List Bullet",
            )
    if m.not_found:
        missing_heading = doc.add_paragraph()
        missing_heading.add_run("Not found in source document: ").bold = True
        missing_heading.add_run(", ".join(m.not_found))

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = foot.add_run(
        f"Pipeline cost ${run.total_cost_usd:.4f} · {run.total_latency_ms / 1000:.1f}s · "
        f"{len(run.telemetry)} model calls"
    )
    fr.font.size = Pt(8)
    fr.font.color.rgb = GREY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
